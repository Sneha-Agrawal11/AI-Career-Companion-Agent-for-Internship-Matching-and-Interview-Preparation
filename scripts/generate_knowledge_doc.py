import os
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

def add_heading(doc, text, level=1):
    heading = doc.add_heading(text, level=level)
    return heading

def add_paragraph(doc, text):
    p = doc.add_paragraph(text)
    return p

def main():
    doc = Document()
    
    # Title
    title = doc.add_heading('InternMatch Product Knowledge Document', 0)
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    
    # 1. Product Overview
    add_heading(doc, '1. Product Overview', level=1)
    
    add_heading(doc, 'Product Name', level=2)
    add_paragraph(doc, 'InternMatch')
    
    add_heading(doc, 'Product Description', level=2)
    add_paragraph(doc, 'InternMatch is an AI-powered career companion platform designed to intelligently match students with internship opportunities based on their resumes, analyze their skill gaps, and help them draft tailored cover letters.')
    
    add_heading(doc, 'Problem it Solves', level=2)
    add_paragraph(doc, 'Students often struggle to find internships that align with their actual skills, and they lack clear insights into what skills they are missing for their desired roles. InternMatch automates resume parsing and compares candidate profiles against internship requirements to provide actionable skill gap analysis and application workflows.')
    
    add_heading(doc, 'Target Users', level=2)
    add_paragraph(doc, 'University students, recent graduates, and entry-level job seekers looking for internships.')
    
    add_heading(doc, 'Main Objectives', level=2)
    add_paragraph(doc, '1. Parse and extract structured data from user resumes using AI.\n2. Match users to relevant internship opportunities using semantic similarity and skill matching.\n3. Highlight skill gaps to guide user learning.\n4. Assist users in generating tailored cover letters for specific applications.')
    
    # 2. Product Features and Functionalities
    add_heading(doc, '2. Product Features and Functionalities', level=1)
    
    features = [
        {
            "name": "Authentication (Login / Register)",
            "purpose": "Secure user access and separate user profiles.",
            "works": "Uses JWT-based authentication to manage sessions. Users can register and login to access their personalized dashboard.",
            "use": "Users enter their email and password on the landing page to authenticate.",
            "result": "Users gain access to the authenticated workspace."
        },
        {
            "name": "Resume Management and Parsing",
            "purpose": "Allow users to store multiple resumes and extract their data.",
            "works": "Users upload PDF resumes. The backend extracts text using PDFPlumber/PyPDF2 and parses it into structured candidate data (skills, experience, education) using an LLM (Groq) or Regex fallback.",
            "use": "Users drag and drop a PDF on the dashboard or upload a new resume via the 'Upload New Resume' button.",
            "result": "Structured profile data is extracted and the resume is saved in the user's Resume Library."
        },
        {
            "name": "Internship Matching",
            "purpose": "Find the best internship opportunities for the user's resume.",
            "works": "Uses FAISS to perform semantic similarity search against a database of internships, and calculates a final score based on skill, education, and experience match percentages.",
            "use": "After selecting an active resume, users click 'Find Internships' to view matching opportunities.",
            "result": "A ranked list of matching internships with match scores."
        },
        {
            "name": "Skill Gap Analysis",
            "purpose": "Identify missing skills required for a specific internship.",
            "works": "Compares the parsed skills from the user's active resume against the required skills of the selected internship.",
            "use": "Users click 'Review Skill Gap' on an internship match to see a breakdown of matched skills vs. missing skills.",
            "result": "Visual breakdown of 'What you already have' vs 'What you are missing'."
        },
        {
            "name": "Cover Letter Generation",
            "purpose": "Draft a personalized cover letter for an internship.",
            "works": "Combines the user's parsed resume data and the internship details, and prompts an LLM to draft a professional cover letter.",
            "use": "Users click 'Draft Cover Letter' from the Skill Gap or Match page.",
            "result": "An editable text area containing the AI-generated cover letter, which can be copied."
        },
        {
            "name": "AI Product Assistant (Chatbot)",
            "purpose": "Answer user questions about the InternMatch platform.",
            "works": "A RAG-based chatbot that retrieves documentation from a ChromaDB vector store and uses Groq to generate grounded responses while remembering conversation history.",
            "use": "Users click the floating AI robot icon in the bottom right corner of the workspace and type questions.",
            "result": "The chatbot answers questions specifically about how to use InternMatch."
        },
        {
            "name": "Theme Toggle (Light/Dark Mode)",
            "purpose": "Provide visual comfort based on user preference.",
            "works": "Toggles a data-theme attribute on the root HTML element, applying specific CSS variable overrides.",
            "use": "Users click the Sun/Moon icon at the bottom of the sidebar.",
            "result": "The entire UI seamlessly switches between light and dark visual themes."
        }
    ]
    
    for f in features:
        add_heading(doc, f['name'], level=2)
        add_paragraph(doc, f'Purpose: {f["purpose"]}')
        add_paragraph(doc, f'How it works: {f["works"]}')
        add_paragraph(doc, f'How users use it: {f["use"]}')
        add_paragraph(doc, f'Expected result: {f["result"]}')
    
    # 3. How to Use the Product
    add_heading(doc, '3. How to Use the Product', level=1)
    
    add_heading(doc, 'Creating an Account & Logging In', level=2)
    add_paragraph(doc, '1. Navigate to the landing page.\n2. Click "Create Account" and provide a full name, email, and password.\n3. Upon successful registration, the user is redirected to the login view.\n4. Enter credentials and click "Sign In" to access the dashboard.')
    
    add_heading(doc, 'Uploading and Parsing a Resume', level=2)
    add_paragraph(doc, '1. In the Dashboard, drag and drop a PDF resume into the upload zone, or click to browse files.\n2. The system analyzes the resume and displays an extraction success screen.\n3. The user can review the parsed skills, experience, and education.\n4. Click "Find Internships" to proceed.')
    
    add_heading(doc, 'Viewing Matches & Skill Gaps', level=2)
    add_paragraph(doc, '1. The Matches page displays a grid of internships ranked by their match score.\n2. Click on a specific match to view details.\n3. Click "Review Skill Gap" to see which required skills are present on the resume and which are missing.')
    
    add_heading(doc, 'Generating a Cover Letter', level=2)
    add_paragraph(doc, '1. From the Skill Gap or Match page, click the option to Draft Cover Letter.\n2. Wait for the AI to generate the draft.\n3. Review, manually personalize if needed, and click "Copy Cover Letter" to use it in an external application.')
    
    add_heading(doc, 'Using the AI Product Assistant', level=2)
    add_paragraph(doc, '1. Click the floating AI robot button on the bottom right of the screen.\n2. Type a question about InternMatch, such as "How do I upload a resume?".\n3. The assistant will provide a helpful answer based on product documentation.')
    
    # 4. Product Architecture
    add_heading(doc, '4. Product Architecture', level=1)
    add_paragraph(doc, 'InternMatch uses a modern client-server architecture with AI integration.')
    
    add_paragraph(doc, 'Frontend:\n- Built with React and Vite.\n- Manages state for authentication, theming, and active resumes.\n- Communicates with the backend via REST APIs.')
    
    add_paragraph(doc, 'Backend/API:\n- Built with FastAPI (Python).\n- Provides endpoints for auth, user management, resume parsing, internship matching, and chat sessions.\n- Orchestrates calls to AI services.')
    
    add_paragraph(doc, 'Application Database:\n- Uses SQLite via SQLAlchemy ORM.\n- Stores Users, Resumes (with parsed JSON data), ChatSessions, and ChatMessages.')
    
    add_paragraph(doc, 'AI/RAG Service & Vector Databases:\n- Internship Matching uses FAISS for semantic similarity between resume skills and internship requirements.\n- The AI Product Assistant uses ChromaDB for retrieving product documentation chunks (RAG).\n- SentenceTransformers ("all-MiniLM-L6-v2") generates text embeddings.')
    
    add_paragraph(doc, 'LLM:\n- Uses Groq API (e.g., Llama 3) for resume parsing, cover letter generation, and Chatbot response generation.')
    
    add_paragraph(doc, 'Authentication:\n- Uses OAuth2 with Password Flow and JWT tokens (python-jose & passlib).')
    
    # 5. Technology Stack
    add_heading(doc, '5. Technology Stack', level=1)
    
    techs = [
        ("React", "Frontend UI library used for building interactive, component-based user interfaces."),
        ("FastAPI", "High-performance Python web framework used for the backend REST API."),
        ("SQLAlchemy & SQLite", "ORM and lightweight database used for persistent application state (users, resumes, chat memory)."),
        ("Groq API", "Provides fast LLM inference for AI features (parsing, cover letters, chatbot)."),
        ("ChromaDB", "Vector database used exclusively for the Chatbot RAG pipeline to store and retrieve product documentation embeddings."),
        ("FAISS", "In-memory vector similarity search library used specifically for rapid internship-to-resume matching."),
        ("SentenceTransformers", "Generates high-quality dense vector embeddings for text chunks and queries locally without API latency."),
        ("PDFPlumber / PyPDF2", "Used to extract raw text from uploaded user PDF resumes before LLM parsing.")
    ]
    
    for name, reason in techs:
        add_paragraph(doc, f'{name}: {reason}')
        
    doc.save('InternMatch_Product_Knowledge.docx')
    print("Generated InternMatch_Product_Knowledge.docx")

if __name__ == '__main__':
    main()
